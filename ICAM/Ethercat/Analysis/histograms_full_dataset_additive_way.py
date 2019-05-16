import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from operator import add
import docx


# Declare functions
def dataframe_to_word(df, filename):
    # open an existing document
    doc = docx.Document(r'C:\Users\exprivia\PycharmProjects\untitled\ICAM\Ethercat\Analysis\\'+filename+'.docx')
    # add a table to the end and create a reference variable
    # extra row is so we can add the header row
    t = doc.add_table(df.shape[0] + 1, df.shape[1])
    # add the header rows.
    for j in range(df.shape[-1]):
        t.cell(0, j).text = df.columns[j]
    # add the rest of the data frame
    for i in range(df.shape[0]):
        for j in range(df.shape[-1]):
            t.cell(i + 1, j).text = str(df.values[i, j])
    # save the doc
    doc.save(r'C:\Users\exprivia\PycharmProjects\untitled\ICAM\Ethercat\Analysis\\'+filename+'.docx')


def to_frequency_dataframe(bins2, frequencies):
    columns = ['edge1', 'edge2', 'abs_frequency']
    df_frequencies = []
    for j in range(len(bins2)):
        print(j)
        if j == len(bins2)-2:
            interval = bins2[j:]
            interval.append(frequencies[j])
            print(interval)
            df_frequencies.append(interval)
            break
        interval = bins2[j:j + 2]
        interval.append(frequencies[j])
        print(interval)
        df_frequencies.append(interval)
    df_frequencies = pd.DataFrame(data=df_frequencies, columns=columns)
    return df_frequencies


# Declare variables
statistics_spread = pd.read_csv('chunks_statistics_spread.csv')
min_spread = statistics_spread['min'].values.min()
max_spread = statistics_spread['max'].values.max()

bins = np.linspace(min_spread, max_spread, num=200, endpoint=True)

chunks_directory = r'..\Data\chunks\ethercat_icam_v7_chunk'

Z_frequencies = [0] * 199
Y_frequencies = [0] * 199
X_frequencies = [0] * 199
n_rows = 0
sumZ = 0
sumY = 0
sumX = 0

# These variables are not used because I will use them to build the movement model (TODO)
z_y_x = pd.DataFrame(columns=['is_moving_Z', 'is_moving_Y', 'new_df_X'])
z_y_x = pd.DataFrame(columns=['zfrequencymonitor', 'yvelocityfeedbackvalue', 'xvelocityfeedbackvalue'])


# Loop over the chunks (from 3 to 9) to calculate the frequencies for each bin
for i in range(2, 9):
    df = pd.read_csv(chunks_directory+str(i+1)+'_events.csv')
    df.sort_values(by=['timestamp_'], inplace=True)

    spread1 = (np.array(df['yactualposition'].values) - np.array(df['ypositionfeedback1value'].values)).reshape(
        len(df['ypositionfeedback1value'].values), 1)
    ZFrequencyMonitor = np.array(df['zfrequencymonitor'].values).reshape(len(df['zfrequencymonitor'].values), 1)
    YVelocityFeedbackValue = np.array(df['yvelocityfeedbackvalue'].values).reshape(
        len(df['yvelocityfeedbackvalue'].values), 1)
    XVelocityFeedbackValue = np.array(df['xvelocityfeedbackvalue'].values).reshape(
        len(df['xvelocityfeedbackvalue'].values), 1)

    new_df = pd.DataFrame(
        data=np.concatenate((spread1, ZFrequencyMonitor, YVelocityFeedbackValue, XVelocityFeedbackValue), axis=1),
        columns=['spread1', 'zfrequencymonitor', 'yvelocityfeedbackvalue', 'xvelocityfeedbackvalue'])

    is_moving_Z = new_df['zfrequencymonitor'] != 0
    is_moving_Y = new_df['yvelocityfeedbackvalue'] != 0
    is_moving_X = new_df['xvelocityfeedbackvalue'] != 0

    new_df_Z = new_df[is_moving_Z]
    new_df_Y = new_df[is_moving_Y]
    new_df_X = new_df[is_moving_X]

    sumZ = sumZ + sum(new_df_Z['spread1'].values)
    sumY = sumY + sum(new_df_Y['spread1'].values)
    sumX = sumX + sum(new_df_X['spread1'].values)

    plt.figure(i)
    nZ, binsZ, patchesZ = plt.hist(new_df_Z['spread1'], bins=bins, alpha=0.5, label='moving_in_Z')
    nY, binsY, patchesY = plt.hist(new_df_Y['spread1'], bins=bins, alpha=0.5, label='moving_in_Y')
    nX, binsX, patchesX = plt.hist(new_df_X['spread1'], bins=bins, alpha=0.5, label='moving_in_X')
    plt.grid(axis='y', alpha=0.75)
    plt.xlabel('Spread1 (10^-4 mm)')
    plt.ylabel('Frequency')
    plt.title('Spread1 Histogram (chunk ' + str(i+1) + ')')
    plt.legend(loc='upper right')
    plt.show()

    Z_frequencies = list(map(add, Z_frequencies, nZ))
    Y_frequencies = list(map(add, Y_frequencies, nY))
    X_frequencies = list(map(add, X_frequencies, nX))
    n_rows = n_rows + len(df['ypositionfeedback1value'].values)


# Write to file (.csv and .docx) the frequencies for each bin, for each axes, as a table
bins1 = bins.tolist()
Z_frequencies_table = to_frequency_dataframe(bins1, Z_frequencies)
Z_frequencies_table.to_csv('Z_frequencies_table.csv', index=False)
dataframe_to_word(Z_frequencies_table, 'Z_frequencies_table')

Y_frequencies_table = to_frequency_dataframe(bins1, Y_frequencies)
Y_frequencies_table.to_csv('Y_frequencies_table.csv', index=False)
dataframe_to_word(Y_frequencies_table, 'Y_frequencies_table')

X_frequencies_table = to_frequency_dataframe(bins1, X_frequencies)
X_frequencies_table.to_csv('X_frequencies_table.csv',index=False)
dataframe_to_word(X_frequencies_table, 'X_frequencies_table')

# Plot the whole histogram
Z_frequencies.append(0)
Y_frequencies.append(0)
X_frequencies.append(0)
plt.figure(i+1)
plt.bar(bins, Z_frequencies, width=456.43216081, align='edge', label='moving_in_Z')
plt.bar(bins, Y_frequencies,  width=456.43216081, align='edge', label='moving_in_Y')
plt.bar(bins, X_frequencies, width=456.43216081, align='edge', label='moving_in_X')
plt.grid(axis='y', alpha=0.75)
plt.xlabel('Spread1 (10^-4 mm)')
plt.ylabel('Frequency')
plt.title('Spread1 Histogram (Full Dataset)')
plt.legend(loc='upper right')
plt.show()

# Calculate the mean
mean_Z = sumZ/sum(Z_frequencies)
mean_Y = sumY/sum(Y_frequencies)
mean_X = sumX/sum(X_frequencies)





